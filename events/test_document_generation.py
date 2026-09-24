import os
from io import BytesIO
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from docx import Document
from pypdf import PdfReader

from events.models import Event
from events.services import generate_participant_list


class DocumentGenerationTests(TestCase):
    fixtures = ["events"]

    @patch("events.services.requests.post")
    def test_docxtpl_renders_reference_participants(self, mock_post):
        mock_post.return_value.content = b"%PDF-1.4\nmock response\n"

        event = Event.objects.get(pk=1)
        generate_participant_list(event)

        self.assertEqual(mock_post.call_count, 1)

        submitted_file = mock_post.call_args.kwargs["files"]["files"]
        self.assertEqual(submitted_file[0], "participant_list.docx")

        docx_content = submitted_file[1].getvalue()
        document = Document(BytesIO(docx_content))

        rendered_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )

        self.assertIn(event.title, rendered_text)

        for participant in event.participants.all():
            with self.subTest(participant=participant.pk):
                self.assertIn(participant.last_name, rendered_text)

        self.assertIn("Lorem ipsum", rendered_text)

    @patch("events.services.requests.post")
    def test_admin_returns_generated_pdf_as_download(self, mock_post):
        mock_post.return_value.content = b"%PDF-1.4\nmock response\n"

        admin_user = get_user_model().objects.create_superuser(
            username="document-test-admin",
            email="document-test@example.invalid",
            password="test-password",
        )
        self.client.force_login(admin_user)

        url = reverse(
            "admin:events_event_generate_participant_list",
            args=[1],
        )
        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn(
            "attachment;",
            response["Content-Disposition"],
        )
        self.assertIn(
            "participant_list_event_1.pdf",
            response["Content-Disposition"],
        )
        self.assertTrue(
            b"".join(response.streaming_content).startswith(b"%PDF-")
        )

    @skipUnless(
        os.environ.get("RUN_GOTENBERG_TESTS") == "1",
        "Requires a running local Gotenberg service.",
    )
    def test_live_gotenberg_creates_multipage_pdf(self):
        event = Event.objects.get(pk=1)
        result = generate_participant_list(event)

        self.assertEqual(result.content_type, "application/pdf")
        self.assertTrue(result.content.startswith(b"%PDF-"))

        reader = PdfReader(BytesIO(result.content))
        self.assertGreater(len(reader.pages), 1)

    @skipUnless(
        os.environ.get("RUN_GOTENBERG_TESTS") == "1",
        "Requires a running local Gotenberg service.",
    )
    def test_live_gotenberg_preserves_participant_text_in_pdf(self):
        for event_id, expected_count in ((1, 60), (2, 25)):
            with self.subTest(event_id=event_id):
                event = Event.objects.get(pk=event_id)
                result = generate_participant_list(event)

                self.assertEqual(result.content_type, "application/pdf")
                self.assertTrue(result.content.startswith(b"%PDF-"))

                reader = PdfReader(BytesIO(result.content))
                self.assertGreaterEqual(len(reader.pages), 1)

                if event_id == 1:
                    self.assertGreater(len(reader.pages), 1)

                pdf_text = "\n".join(
                    page.extract_text() or ""
                    for page in reader.pages
                )

                self.assertIn(event.title, pdf_text)
                self.assertEqual(event.participants.count(), expected_count)

                for participant in event.participants.all():
                    with self.subTest(
                            event_id=event_id,
                            participant_id=participant.pk,
                    ):
                        self.assertIn(participant.last_name, pdf_text)
                if event_id == 1:
                    self.assertIn("Lorem ipsum", pdf_text)
                else:
                    self.assertNotIn("Lorem ipsum", pdf_text)
